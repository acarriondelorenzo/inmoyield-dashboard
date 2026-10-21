"""Tesorería desde "BBDD Inversiones Inmoyield - Clean" (Google Sheet, mantenido por Yago).

Lee la pestaña "Caja y Calendario" (exportada a .xlsx con el conector de Drive),
localizando cada bloque por sus ETIQUETAS. Nunca modifica el archivo origen.
Genera el mismo formato (version 2) que pinta template.html.

Uso: python3 treasury_clean.py clean.xlsx salida.json
Imprime "OK ..." o "TREASURY_VALIDATION_FAILED" (+ motivos) y líneas "AVISO ...".
"""
import datetime, json, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl

sys.path.insert(0, __import__("os").path.dirname(__file__))
from treasury_v22 import MESES, norm, find, num, iso, header_map  # utilidades comunes


def eur(x):
    return f"{x:,.0f} €".replace(",", ".")


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if "caja y calendario" in norm(n)), None)
    out = {"version": 2, "fuente": "BBDD Inversiones Inmoyield - Clean", "avisos": []}
    if ws is None:
        return out, ["No se encuentra la pestaña 'Caja y Calendario'."]
    reasons = []

    p = find(ws, "Fecha de referencia")
    ref = ws.cell(p[0], p[1] + 1).value if p else None
    out["fechaReferencia"] = iso(ref)
    refd = ref.date() if isinstance(ref, datetime.datetime) else datetime.date.today()

    # --- Caja por sociedad
    p = find(ws, "Sociedad")
    soc, total = [], None
    if p:
        h = header_map(ws, p[0])
        r = p[0] + 1
        while r <= ws.max_row:
            name = ws.cell(r, p[1]).value
            if name in (None, ""):
                break
            row = {"sociedad": str(name).strip(), "fase": ws.cell(r, h.get("fase", 2)).value,
                   "caja": num(ws.cell(r, h.get("caja total", 5)).value),
                   "capitalInvertido": num(ws.cell(r, h.get("capital invertido", 7)).value),
                   "cajaProyectada90d": num(ws.cell(r, h.get("caja + cobros - pagos 90 d", 9)).value)}
            if norm(name) == "total":
                total = row
                break
            soc.append(row)
            r += 1
    out["sociedades"] = soc
    nombres_soc = {norm(s["sociedad"]) for s in soc}
    out["cajaTotal"] = total["caja"] if total else (sum(s["caja"] or 0 for s in soc) if soc else None)
    out["capitalInvertido"] = total["capitalInvertido"] if total else None
    if out["cajaTotal"] is None:
        reasons.append("No se pudo leer la caja total (cuadro 'Caja por sociedad').")

    # --- Cuadre de capital sin sociedad
    p = find(ws, "Capital vivo cuyo", startswith=True)
    if p:
        dif = next((num(ws.cell(p[0], c).value) for c in range(p[1] + 1, ws.max_column + 1)
                    if num(ws.cell(p[0], c).value) is not None), None)
        dif = round(dif) if dif is not None else None
        out["capitalSinAsignar"] = dif
        if dif and abs(dif) >= 1:
            out["avisos"].append(f"Hay {eur(dif)} de capital vivo con una «SL» que no coincide con ninguna sociedad.")

    # --- Tesorería mes a mes
    p = find(ws, "Mes")
    meses = []
    if p:
        h = header_map(ws, p[0])
        r = p[0] + 1
        while r <= ws.max_row and isinstance(ws.cell(r, p[1]).value, (datetime.date, datetime.datetime)):
            d = ws.cell(r, p[1]).value
            meses.append({"mes": f"{MESES[d.month - 1]} {d.year}", "iso": iso(d)[:7],
                          "cobros": num(ws.cell(r, h.get("cobros comprometidos", 2)).value) or 0,
                          "pagos": num(ws.cell(r, h.get("pagos comprometidos", 3)).value) or 0,
                          "neto": num(ws.cell(r, h.get("neto del mes", 4)).value) or 0,
                          "cajaAcumulada": num(ws.cell(r, h.get("caja acumulada", 5)).value),
                          "cajaConVencimientos": num(ws.cell(r, h.get("caja acum. con vencimientos", 7)).value)})
            r += 1
    if not meses:
        reasons.append("No se pudo leer el bloque 'Tesorería mes a mes'.")
    out["meses"] = meses
    neg = next((m for m in meses if (m["cajaAcumulada"] or 0) < 0), None)
    out["mesDeficit"] = neg["mes"] if neg else None
    out["minCajaAcumulada"] = min((m["cajaAcumulada"] or 0) for m in meses) if meses else None
    if neg:
        out["avisos"].append(f"La caja prudente entra en negativo en {neg['mes']} ({eur(neg['cajaAcumulada'])}); "
                             f"mínimo del horizonte: {eur(out['minCajaAcumulada'])}.")

    p = find(ws, "Vencimientos con fecha fuera", startswith=True)
    if p:
        fuera = next((num(ws.cell(p[0], c).value) for c in range(p[1] + 1, ws.max_column + 1)
                      if num(ws.cell(p[0], c).value) is not None), None)
        if fuera and abs(fuera) >= 1:
            out["avisos"].append(f"Hay {eur(fuera)} de vencimientos de proyectos vivos con «Est. Fin1» caducada o vacía.")

    # --- Pagos y cobros comprometidos
    p = find(ws, "Concepto")
    movs = []
    if p:
        h = header_map(ws, p[0])
        cF, cC, cP, cS, cT, cI, cE, cN = (h.get(k) for k in ("fecha", "concepto", "proyecto", "sociedad", "tipo", "importe", "estado", "notas"))
        r = p[0] + 1
        while r <= ws.max_row:
            first = norm(ws.cell(r, 1).value)
            if first == "total" or first.startswith("vencimientos de proyectos"):
                break
            concepto = ws.cell(r, cC).value if cC else None
            if concepto not in (None, ""):
                imp = num(ws.cell(r, cI).value) if cI else None
                nota = ws.cell(r, cN).value if cN else None
                m = {"fecha": iso(ws.cell(r, cF).value) if cF else None, "concepto": str(concepto).strip(),
                     "proyecto": ws.cell(r, cP).value if cP else None,
                     "sociedad": ws.cell(r, cS).value if cS else None,
                     "tipo": ws.cell(r, cT).value if cT else None, "importe": imp,
                     "estado": ws.cell(r, cE).value if cE else None, "nota": nota}
                if imp is None:
                    est = num(nota)
                    if est is not None:
                        m["importeNota"] = (-1 if norm(m["tipo"]) == "pago" else 1) * abs(est)
                movs.append(m)
            r += 1
    if not movs:
        reasons.append("No se encontró el cuadro 'Pagos y cobros comprometidos'.")
    movs.sort(key=lambda m: (m["fecha"] or "9999", m["concepto"]))
    out["movimientos"] = movs

    # --- Vencimientos de proyectos vivos: salidas "firmes" (Situación actual = Vendido)
    FIRMES = {"vendido"}
    p = find(ws, "Fecha est. salida")
    venc = []
    if p:
        h = header_map(ws, p[0])
        r = p[0] + 1
        while r <= ws.max_row:
            proj = ws.cell(r, h.get("proyecto", 3)).value
            first = norm(ws.cell(r, 1).value)
            if first.startswith("total"):
                break
            if proj not in (None, ""):
                d = ws.cell(r, p[1]).value
                venc.append({"fecha": iso(d) if isinstance(d, (datetime.date, datetime.datetime)) else None,
                             "proyecto": str(proj).strip(),
                             "situacion": ws.cell(r, h.get("situación actual", 5)).value,
                             "total": num(ws.cell(r, h.get("total a cobrar", 8)).value) or 0})
            r += 1
    firmes = [v for v in venc if norm(v["situacion"]) in FIRMES and v["fecha"]]
    out["salidasFirmes"] = firmes
    for m in meses:
        acum = sum(v["total"] for v in firmes if v["fecha"][:7] <= m["iso"])
        m["cajaConFirmes"] = (m["cajaAcumulada"] or 0) + acum
    negf = next((m for m in meses if m["cajaConFirmes"] < 0), None)
    out["mesDeficitConFirmes"] = negf["mes"] if negf else None
    out["minCajaConFirmes"] = min((m["cajaConFirmes"] for m in meses), default=None)

    # --- 90 días: calculado de forma transparente desde los movimientos con fecha
    lim = (refd + datetime.timedelta(days=90)).isoformat()
    en90 = [m for m in movs if m["fecha"] and m["fecha"] <= lim and m["importe"] is not None and norm(m["estado"]) != "pagado"]
    out["cobros90d"] = sum(m["importe"] for m in en90 if m["importe"] > 0)
    out["pagos90d"] = sum(m["importe"] for m in en90 if m["importe"] < 0)
    if out["cajaTotal"] is not None:
        out["cajaProyectada90d"] = out["cajaTotal"] + out["cobros90d"] + out["pagos90d"]
    if total and total.get("cajaProyectada90d") is not None and out.get("cajaProyectada90d") is not None \
            and abs(total["cajaProyectada90d"] - out["cajaProyectada90d"]) >= 1:
        out["avisos"].append(f"La «Caja + cobros - pagos 90 d» del cuadro por sociedad ({eur(total['cajaProyectada90d'])}) "
                             f"no coincide con la suma de los movimientos a 90 días ({eur(out['cajaProyectada90d'])}); "
                             f"suele deberse a movimientos sin sociedad o con un nombre de sociedad que no existe.")

    # --- Avisos de calidad de los movimientos
    sin_fecha = [m for m in movs if not m["fecha"] and m["importe"]]
    if sin_fecha:
        out["avisos"].append("Movimientos sin fecha (no entran en el calendario): " +
                             "; ".join(f"{m['concepto']} ({eur(m['importe'])})" for m in sin_fecha))
    sin_soc = [m for m in movs if m["importe"] and not m["sociedad"]]
    if sin_soc:
        out["avisos"].append("Movimientos sin sociedad asignada: " +
                             "; ".join(f"{m['concepto']} ({eur(m['importe'])})" for m in sin_soc))
    soc_mal = [m for m in movs if m["sociedad"] and nombres_soc and norm(m["sociedad"]) not in nombres_soc]
    if soc_mal:
        out["avisos"].append("Movimientos con una sociedad que no existe en el cuadro: " +
                             "; ".join(f"{m['concepto']} → «{m['sociedad']}»" for m in soc_mal))
    discrep = []
    for m in movs:
        if m["importe"] is not None:
            nv = num(m["nota"])
            if nv is not None and abs(abs(nv) - abs(m["importe"])) >= 1:
                discrep.append(f"{m['concepto']} (importe {eur(m['importe'])}, nota «{m['nota']}»)")
    if discrep:
        out["avisos"].append("Importe y nota no coinciden: " + "; ".join(discrep))
    sin_imp = [m for m in movs if m["importe"] is None]
    if sin_imp:
        out["avisos"].append(f"{len(sin_imp)} movimiento(s) solo en notas, sin importe firme (fuera de la proyección): " +
                             "; ".join(f"{m['concepto']} ({m['nota']})" if m.get("nota") else m["concepto"] for m in sin_imp))
    noconf = [m for m in en90 if m["importe"] > 0 and norm(m["estado"]) in ("previsto", "por confirmar")]
    if noconf:
        out["avisos"].append("Cobros no confirmados incluidos en la proyección a 90 días: " +
                             "; ".join(f"{m['concepto']} ({eur(m['importe'])}, {m['estado']})" for m in noconf))
    return out, reasons


if __name__ == "__main__":
    t, reasons = parse(sys.argv[1])
    if reasons:
        print("TREASURY_VALIDATION_FAILED")
        for r in reasons:
            print(" -", r)
        sys.exit(1)
    json.dump(t, open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    for a in t["avisos"]:
        print("AVISO", a)
    print(f"OK caja={t['cajaTotal']} proyectada90d={t.get('cajaProyectada90d')} meses={len(t['meses'])} "
          f"movimientos={len(t['movimientos'])} ref={t.get('fechaReferencia')}")
