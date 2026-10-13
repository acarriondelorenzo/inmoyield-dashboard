"""Tesorería desde "BBDD Inversiones Inmoyield - v2.2" (xlsx en Drive, mantenido por Yago).

Lee las pestañas "2. Societario y Caja" y "3. Calendario de Pagos" localizando los
bloques por sus ETIQUETAS (no por coordenadas fijas), para sobrevivir a filas
insertadas. Nunca modifica el archivo origen.

Uso: python3 treasury_v22.py v22.xlsx salida.json
Imprime "OK ..." o "TREASURY_VALIDATION_FAILED" (+ motivos) y avisos "AVISO ...".
"""
import datetime, json, re, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def norm(v):
    return " ".join(str(v).split()).strip().lower() if v is not None else ""


def find(ws, label, startswith=False):
    """Primera celda cuyo texto coincide con label -> (fila, columna)."""
    t = norm(label)
    for row in ws.iter_rows():
        for c in row:
            n = norm(c.value)
            if n == t or (startswith and n.startswith(t)):
                return c.row, c.column
    return None


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").replace(" ", "").strip()
    m = re.fullmatch(r"(-?)([\d.,]+)\s*([kKmM]?)", s)
    if not m:
        return None
    body = m.group(2)
    if "," in body and "." in body:
        body = body.replace(".", "").replace(",", ".")
    elif "," in body:
        body = body.replace(",", ".")
    try:
        x = float(body)
    except ValueError:
        return None
    mult = {"k": 1e3, "m": 1e6}.get(m.group(3).lower(), 1)
    return (-1 if m.group(1) else 1) * x * mult


def eur(x):
    return f"{x:,.0f} €".replace(",", ".")


def iso(d):
    if isinstance(d, datetime.datetime):
        d = d.date()
    return d.isoformat() if isinstance(d, datetime.date) else None


def header_map(ws, row):
    return {norm(ws.cell(row, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value not in (None, "")}


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    names = {norm(n): n for n in wb.sheetnames}
    sc = next((wb[names[k]] for k in names if "societario y caja" in k), None)
    cal = next((wb[names[k]] for k in names if "calendario de pagos" in k), None)
    out = {"version": 2, "fuente": "BBDD Inversiones Inmoyield - v2.2", "avisos": []}
    reasons = []
    if sc is None or cal is None:
        return out, ["No se encuentran las pestañas 'Societario y Caja' y/o 'Calendario de Pagos'."]

    # --- Fecha de referencia
    p = find(sc, "Fecha de referencia")
    if p:
        out["fechaReferencia"] = iso(sc.cell(p[0], p[1] + 1).value)

    # --- Indicadores clave (etiquetas en una fila, valores en la de debajo)
    kp = {"cajaTotal": "Caja total disponible", "capitalInvertido": "Capital invertido (vivo)",
          "bEstimadoPendiente": "B° estimado pendiente", "cobros90d": "Cobros previstos 90 d",
          "pagos90d": "Pagos comprometidos 90 d", "cajaProyectada90d": "Caja proyectada 90 d"}
    for k, lab in kp.items():
        p = find(sc, lab)
        out[k] = num(sc.cell(p[0] + 1, p[1]).value) if p else None
        if out[k] is None:
            reasons.append(f"No se pudo leer '{lab}' en 'Societario y Caja'.")

    # --- Caja por sociedad
    p = find(sc, "Sociedad")
    soc = []
    if p:
        h = header_map(sc, p[0])
        r = p[0] + 1
        while r <= sc.max_row:
            name = sc.cell(r, p[1]).value
            if name in (None, "") or norm(name) == "total":
                break
            soc.append({"sociedad": str(name).strip(), "fase": sc.cell(r, h.get("fase", 2)).value,
                        "caja": num(sc.cell(r, h.get("caja total", 5)).value),
                        "capitalInvertido": num(sc.cell(r, h.get("capital invertido", 7)).value),
                        "cajaProyectada90d": num(sc.cell(r, h.get("caja proyectada 90 d", 11)).value)})
            r += 1
    out["sociedades"] = soc

    # --- Cuadre contra la BBDD
    p = find(sc, "Diferencia sin asignar", startswith=True)
    if p:
        dif = num(next((sc.cell(p[0], c).value for c in range(p[1] + 1, sc.max_column + 1)
                        if num(sc.cell(p[0], c).value) is not None), None))
        out["capitalSinAsignar"] = dif
        if dif:
            out["avisos"].append(f"Hay {eur(dif)} de capital invertido sin sociedad asignada (cuadre de 'Societario y Caja' distinto de 0).")

    # --- Resumen mensual
    p = find(cal, "Mes")
    meses = []
    if p:
        h = header_map(cal, p[0])
        r = p[0] + 1
        while r <= cal.max_row and isinstance(cal.cell(r, p[1]).value, (datetime.date, datetime.datetime)):
            d = cal.cell(r, p[1]).value
            meses.append({"mes": f"{MESES[d.month - 1]} {d.year}", "iso": iso(d)[:7],
                          "cobros": num(cal.cell(r, h.get("cobros comprometidos", 2)).value) or 0,
                          "pagos": num(cal.cell(r, h.get("pagos comprometidos", 3)).value) or 0,
                          "neto": num(cal.cell(r, h.get("neto del mes", 4)).value) or 0,
                          "cajaAcumulada": num(cal.cell(r, h.get("caja acumulada", 5)).value),
                          "cajaConVencimientos": num(cal.cell(r, h.get("caja acumulada con vencimientos", 7)).value)})
            r += 1
    if not meses:
        reasons.append("No se pudo leer el resumen mensual de 'Calendario de Pagos'.")
    out["meses"] = meses
    neg = next((m for m in meses if (m["cajaAcumulada"] or 0) < 0), None)
    out["mesDeficit"] = neg["mes"] if neg else None

    # --- Pagos y cobros comprometidos (cuadro editable)
    p = find(cal, "Concepto")
    movs = []
    if p:
        h = header_map(cal, p[0])
        cF, cC, cP, cS, cT, cI, cE, cN = (h.get(k) for k in ("fecha", "concepto", "proyecto", "sociedad", "tipo", "importe", "estado", "notas y conciliación"))
        r = p[0] + 1
        while r <= cal.max_row:
            first = norm(cal.cell(r, 1).value)
            if first == "total" or "vencimientos estimados" in first:
                break
            concepto = cal.cell(r, cC).value if cC else None
            if concepto not in (None, ""):
                imp = num(cal.cell(r, cI).value) if cI else None
                nota = cal.cell(r, cN).value if cN else None
                m = {"fecha": iso(cal.cell(r, cF).value) if cF else None, "concepto": str(concepto).strip(),
                     "proyecto": cal.cell(r, cP).value if cP else None, "sociedad": cal.cell(r, cS).value if cS else None,
                     "tipo": cal.cell(r, cT).value if cT else None, "importe": imp,
                     "estado": cal.cell(r, cE).value if cE else None, "nota": nota}
                if imp is None:
                    est = num(nota)
                    if est is not None:
                        signo = -1 if norm(m["tipo"]) == "pago" else 1
                        m["importeNota"] = signo * abs(est)
                movs.append(m)
            r += 1
    movs.sort(key=lambda m: (m["fecha"] or "9999", m["concepto"]))
    out["movimientos"] = movs
    sin = [m for m in movs if m["importe"] is None]
    if sin:
        tot = sum(m.get("importeNota", 0) for m in sin)
        out["sinImporteNeto"] = tot
        out["avisos"].append(
            f"{len(sin)} movimientos sin cifra en la columna Importe (no entran en la proyección): "
            + "; ".join(f"{m['concepto']} ({m['nota']})" if m.get("nota") else m["concepto"] for m in sin)
            + (f". Neto según notas: {eur(tot)}" if tot else ""))
    noconf = [m for m in movs if m["importe"] and m["importe"] > 0 and norm(m["estado"]) in ("previsto", "por confirmar")]
    if noconf:
        out["avisos"].append("Cobros no confirmados incluidos en la proyección: " +
                             "; ".join(f"{m['concepto']} ({eur(m['importe'])}, {m['estado']})" for m in noconf))
    if not movs:
        reasons.append("No se encontró el cuadro de pagos y cobros comprometidos.")
    return out, reasons


if __name__ == "__main__":
    t, reasons = parse(sys.argv[1])
    if reasons:
        print("TREASURY_VALIDATION_FAILED")
        for r in reasons:
            print(" -", r)
        sys.exit(1)
    json.dump(t, open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    for a in t["avisos"]:
        print("AVISO", a)
    print(f"OK caja={t['cajaTotal']} proyectada90d={t['cajaProyectada90d']} meses={len(t['meses'])} movimientos={len(t['movimientos'])} ref={t.get('fechaReferencia')}")
